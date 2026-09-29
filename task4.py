import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from pathlib import Path
from scipy.stats import ttest_ind, chi2_contingency
from pptx import Presentation
from pptx.util import Inches


# ==========================================
# TASK 4: DATA STORYTELLING & STATISTICAL VALIDATION
# ==========================================

print("=" * 60)
print("TASK 4: DATA STORYTELLING & STATISTICAL VALIDATION")
print("=" * 60)


# ==========================================
# PART 1: CREATE OUTPUT FOLDER
# ==========================================

output_folder = Path("task4_outputs")
output_folder.mkdir(exist_ok=True)


# ==========================================
# PART 2: LOAD DATASET
# ==========================================

file_path = Path("sales.csv")

if not file_path.exists():
    file_path = Path("cleaned_sales.csv")

if not file_path.exists():
    raise FileNotFoundError(
        "sales.csv or cleaned_sales.csv not found!"
    )

df = pd.read_csv(file_path)

print("\nDataset Loaded Successfully!")
print("Rows:", len(df))
print("Columns:", list(df.columns))


# ==========================================
# PART 3: CLEAN COLUMN NAMES
# ==========================================

df.columns = (
    df.columns
    .str.strip()
    .str.lower()
    .str.replace(" ", "_")
)


# ==========================================
# PART 4: IDENTIFY IMPORTANT COLUMNS
# ==========================================

def find_column(possible_names):

    for name in possible_names:

        if name in df.columns:
            return name

    return None


sales_col = find_column([
    "sales",
    "total_sales",
    "revenue",
    "amount",
    "total_amount"
])

category_col = find_column([
    "category",
    "product_category"
])

product_col = find_column([
    "product",
    "product_name",
    "item"
])

quantity_col = find_column([
    "quantity",
    "qty",
    "units_sold"
])


if sales_col is None:

    if "price" in df.columns and quantity_col is not None:

        df["calculated_sales"] = (
            pd.to_numeric(df["price"], errors="coerce")
            * pd.to_numeric(df[quantity_col], errors="coerce")
        )

        sales_col = "calculated_sales"

    else:

        raise ValueError(
            "Sales column not found. Check your dataset."
        )


df[sales_col] = pd.to_numeric(
    df[sales_col],
    errors="coerce"
)

df = df.dropna(subset=[sales_col])


print("\nSelected Sales Column:", sales_col)


# ==========================================
# PART 5: BASIC DATA SUMMARY
# ==========================================

total_sales = df[sales_col].sum()

average_sales = df[sales_col].mean()

minimum_sales = df[sales_col].min()

maximum_sales = df[sales_col].max()

total_records = len(df)


print("\n" + "=" * 60)
print("DATA SUMMARY")
print("=" * 60)

print("Total Sales:", round(total_sales, 2))

print("Average Sales:", round(average_sales, 2))

print("Minimum Sales:", round(minimum_sales, 2))

print("Maximum Sales:", round(maximum_sales, 2))

print("Total Records:", total_records)


# ==========================================
# PART 6: CATEGORY ANALYSIS
# ==========================================

category_summary = None

if category_col is not None:

    category_summary = (
        df.groupby(category_col)[sales_col]
        .agg(["sum", "mean", "count"])
        .reset_index()
    )

    category_summary = category_summary.sort_values(
        "sum",
        ascending=False
    )

    category_summary.to_csv(
        output_folder / "category_summary.csv",
        index=False
    )

    print("\nCATEGORY ANALYSIS")

    print(category_summary)

else:

    print("\nCategory column not found.")


# ==========================================
# PART 7: PRODUCT ANALYSIS
# ==========================================

product_summary = None

if product_col is not None:

    product_summary = (
        df.groupby(product_col)[sales_col]
        .sum()
        .reset_index()
        .sort_values(sales_col, ascending=False)
    )

    product_summary.to_csv(
        output_folder / "product_summary.csv",
        index=False
    )

    print("\nPRODUCT ANALYSIS")

    print(product_summary)


# ==========================================
# PART 8: CREATE SALES CHART
# ==========================================

plt.figure(figsize=(10, 6))

if category_summary is not None:

    plt.bar(
        category_summary[category_col].astype(str),
        category_summary["sum"]
    )

    plt.title("Sales by Category")

    plt.xlabel("Category")

    plt.ylabel("Total Sales")

    plt.xticks(rotation=30)

else:

    plt.hist(df[sales_col], bins=10)

    plt.title("Sales Distribution")

    plt.xlabel("Sales")

    plt.ylabel("Frequency")


plt.tight_layout()

plt.savefig(
    output_folder / "sales_analysis.png",
    dpi=150
)

plt.close()


# ==========================================
# PART 9: DATA STORY
# ==========================================

story_lines = []

story_lines.append("TASK 4: DATA STORYTELLING")
story_lines.append("=" * 50)

story_lines.append(
    f"Total sales recorded: {total_sales:.2f}"
)

story_lines.append(
    f"Average sales value: {average_sales:.2f}"
)

story_lines.append(
    f"Number of records: {total_records}"
)

story_lines.append(
    f"Minimum sales value: {minimum_sales:.2f}"
)

story_lines.append(
    f"Maximum sales value: {maximum_sales:.2f}"
)


if category_summary is not None:

    top_category = category_summary.iloc[0][category_col]

    top_category_sales = category_summary.iloc[0]["sum"]

    story_lines.append(
        f"Highest sales category: {top_category}"
    )

    story_lines.append(
        f"Highest category sales: {top_category_sales:.2f}"
    )


if product_summary is not None:

    top_product = product_summary.iloc[0][product_col]

    top_product_sales = product_summary.iloc[0][sales_col]

    story_lines.append(
        f"Highest sales product: {top_product}"
    )

    story_lines.append(
        f"Highest product sales: {top_product_sales:.2f}"
    )


story_lines.append("")
story_lines.append("BUSINESS OBSERVATIONS")
story_lines.append(
    "1. Sales performance can be compared across categories."
)
story_lines.append(
    "2. Product-level sales help identify revenue contributors."
)
story_lines.append(
    "3. Statistical testing is used to examine differences "
    "and relationships in the dataset."
)


story_text = "\n".join(story_lines)

print("\n" + story_text)

with open(
    output_folder / "data_story.txt",
    "w",
    encoding="utf-8"
) as file:

    file.write(story_text)


# ==========================================
# PART 10: T-TEST
# ==========================================

print("\n" + "=" * 60)
print("HYPOTHESIS TESTING: T-TEST")
print("=" * 60)

ttest_result = "T-test not performed."

if category_col is not None:

    groups = []

    for category, group in df.groupby(category_col):

        values = group[sales_col].dropna()

        if len(values) >= 2:

            groups.append((category, values))

    if len(groups) >= 2:

        group1_name, group1 = groups[0]

        group2_name, group2 = groups[1]

        t_stat, p_value = ttest_ind(
            group1,
            group2,
            equal_var=False
        )

        print("Group 1:", group1_name)

        print("Group 2:", group2_name)

        print("T-statistic:", round(t_stat, 4))

        print("P-value:", round(p_value, 4))

        alpha = 0.05

        if p_value < alpha:

            conclusion = (
                "Reject the null hypothesis. "
                "The groups show a statistically significant "
                "difference in sales."
            )

        else:

            conclusion = (
                "Fail to reject the null hypothesis. "
                "There is insufficient evidence of a "
                "statistically significant difference."
            )

        print(conclusion)

        ttest_result = (
            f"T-Test comparing {group1_name} and {group2_name}\n"
            f"T-statistic: {t_stat:.4f}\n"
            f"P-value: {p_value:.4f}\n"
            f"Conclusion: {conclusion}"
        )

    else:

        print(
            "At least two categories with enough data "
            "are required for the t-test."
        )

else:

    print("Category column not found. T-test skipped.")


# ==========================================
# PART 11: CHI-SQUARE TEST
# ==========================================

print("\n" + "=" * 60)
print("HYPOTHESIS TESTING: CHI-SQUARE TEST")
print("=" * 60)

chi_result = "Chi-square test not performed."

if category_col is not None and product_col is not None:

    contingency_table = pd.crosstab(
        df[category_col],
        df[product_col]
    )

    if (
        contingency_table.shape[0] >= 2
        and contingency_table.shape[1] >= 2
    ):

        chi_stat, chi_p, dof, expected = chi2_contingency(
            contingency_table
        )

        print("Chi-square statistic:", round(chi_stat, 4))

        print("P-value:", round(chi_p, 4))

        print("Degrees of freedom:", dof)

        alpha = 0.05

        if chi_p < alpha:

            chi_conclusion = (
                "Reject the null hypothesis. "
                "There is evidence of an association "
                "between category and product."
            )

        else:

            chi_conclusion = (
                "Fail to reject the null hypothesis. "
                "There is insufficient evidence of "
                "an association between category and product."
            )

        print(chi_conclusion)

        chi_result = (
            "Chi-Square Test\n"
            f"Chi-square statistic: {chi_stat:.4f}\n"
            f"P-value: {chi_p:.4f}\n"
            f"Degrees of freedom: {dof}\n"
            f"Conclusion: {chi_conclusion}"
        )

    else:

        print(
            "Chi-square test requires at least "
            "two categories and two products."
        )

else:

    print(
        "Category or product column not found. "
        "Chi-square test skipped."
    )


# ==========================================
# PART 12: SAVE STATISTICAL REPORT
# ==========================================

statistical_report = (
    "TASK 4: STATISTICAL VALIDATION\n"
    + "=" * 50
    + "\n\n"
    + ttest_result
    + "\n\n"
    + chi_result
)

with open(
    output_folder / "statistical_report.txt",
    "w",
    encoding="utf-8"
) as file:

    file.write(statistical_report)


# ==========================================
# PART 13: CREATE POWERPOINT
# ==========================================

print("\nCreating PowerPoint Presentation...")

presentation = Presentation()


# Slide 1: Title

slide = presentation.slides.add_slide(
    presentation.slide_layouts[0]
)

slide.shapes.title.text = (
    "Sales Data Storytelling & Statistical Validation"
)

slide.placeholders[1].text = (
    "Task 4 - Data Analytics Internship\n"
    "Prepared by Neha Patra"
)


# Slide 2: Data Summary

slide = presentation.slides.add_slide(
    presentation.slide_layouts[1]
)

slide.shapes.title.text = "Data Summary"

slide.placeholders[1].text = (
    f"Total Sales: {total_sales:.2f}\n"
    f"Average Sales: {average_sales:.2f}\n"
    f"Total Records: {total_records}\n"
    f"Minimum Sales: {minimum_sales:.2f}\n"
    f"Maximum Sales: {maximum_sales:.2f}"
)


# Slide 3: Business Insights

slide = presentation.slides.add_slide(
    presentation.slide_layouts[1]
)

slide.shapes.title.text = "Business Insights"

slide.placeholders[1].text = (
    "1. Category-wise sales comparison\n"
    "2. Product-level sales analysis\n"
    "3. Identification of key sales contributors\n"
    "4. Statistical validation of business patterns"
)


# Slide 4: Statistical Testing

slide = presentation.slides.add_slide(
    presentation.slide_layouts[1]
)

slide.shapes.title.text = "Statistical Validation"

slide.placeholders[1].text = (
    ttest_result
    + "\n\n"
    + chi_result
)


# Slide 5: Conclusion

slide = presentation.slides.add_slide(
    presentation.slide_layouts[1]
)

slide.shapes.title.text = "Conclusion"

slide.placeholders[1].text = (
    "The analysis summarizes sales performance, "
    "compares business groups, and applies statistical "
    "tests to examine differences and relationships.\n\n"
    "Further analysis can help support business decisions."
)


presentation.save(
    output_folder / "task4_presentation.pptx"
)


# ==========================================
# FINAL OUTPUT
# ==========================================

print("\n" + "=" * 60)

print("TASK 4 COMPLETED SUCCESSFULLY")

print("=" * 60)

print("\nFiles saved inside:")

print("task4_outputs")

print("\nGenerated Files:")

print("1. category_summary.csv")

print("2. product_summary.csv")

print("3. sales_analysis.png")

print("4. data_story.txt")

print("5. statistical_report.txt")

print("6. task4_presentation.pptx")

print("\nThank You!")